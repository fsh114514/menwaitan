#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 docs/*.md 转成 .docx（赛事要求 AI 使用说明交 Word 格式）。
只用 python-docx，支持：标题 / 段落 / 表格 / 列表 / 粗体 / 行内代码 / 分隔线。
"""
import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

SERIF = "宋体"
SANS = "黑体"
INK = RGBColor(0x17, 0x12, 0x0D)
RED = RGBColor(0xA3, 0x27, 0x2C)
SOFT = RGBColor(0x5A, 0x51, 0x47)


def set_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_rich(par, text, base_name=SERIF, size=10.5, color=INK):
    """处理 **粗体** 与 `代码`，其余按普通文本。"""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tk in tokens:
        if not tk:
            continue
        if tk.startswith("**") and tk.endswith("**"):
            r = par.add_run(tk[2:-2])
            set_font(r, SANS, size, bold=True, color=color)
        elif tk.startswith("`") and tk.endswith("`"):
            r = par.add_run(tk[1:-1])
            set_font(r, "Menlo", size - 1, color=RED)
        else:
            r = par.add_run(tk)
            set_font(r, base_name, size, color=color)


def convert(md_path, docx_path, title=None):
    md = io.open(md_path, encoding="utf-8").read()
    lines = md.split("\n")
    doc = Document()

    # 页面与正文默认样式
    for sec in doc.sections:
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.4)
        sec.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)

    i = 0
    n = len(lines)
    first_h1 = True
    while i < n:
        raw = lines[i]
        line = raw.rstrip()

        # 空行
        if not line.strip():
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}", line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            from docx.oxml import OxmlElement
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D8D0C0")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # 表格：连续的 | 行
        if line.lstrip().startswith("|"):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            # 去掉 |---|
            rows = [r for r in rows if not re.fullmatch(r"\|[\s:\-|]+\|", r)]

            def cells(r):
                return [c.strip() for c in r.strip().strip("|").split("|")]

            if not rows:
                continue
            body = [cells(r) for r in rows]
            ncol = max(len(r) for r in body)
            table = doc.add_table(rows=len(body), cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(body):
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ""
                    cell = table.cell(ri, ci)
                    cell.paragraphs[0].text = ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_before = Pt(2)
                    par.paragraph_format.space_after = Pt(2)
                    add_rich(par, txt, size=9.5,
                             color=(INK if ri else RED))
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                set_font(r, SANS, 19, bold=True, color=INK)
                p.paragraph_format.space_after = Pt(4)
                # 红规
                pPr = p._p.get_or_add_pPr()
                from docx.oxml import OxmlElement
                pbdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "18")
                bottom.set(qn("w:color"), "A3272C")
                pbdr.append(bottom)
                pPr.append(pbdr)
                first_h1 = False
            else:
                size = {2: 14.5, 3: 12.5, 4: 11}[level]
                p = doc.add_paragraph()
                r = p.add_run(text)
                set_font(r, SANS, size, bold=True,
                         color=(INK if level <= 2 else RED))
                p.paragraph_format.space_before = Pt(12 if level == 2 else 8)
                p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1))
            text = m.group(2)
            style = "List Bullet" if indent < 2 else "List Bullet 2"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            add_rich(p, text)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 引用块
        if line.startswith(">"):
            text = line.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich(p, text, size=9.5, color=SOFT)
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_rich(p, m.group(3))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_rich(p, line)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.45
        i += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "docs/AI工具使用说明.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "docs/AI工具使用说明.docx"
    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    convert(src, dst)
    mb = os.path.getsize(dst) / 1048576
    print(f"  ✅ {src} → {dst} ({mb:.2f} MB)")
